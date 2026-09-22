"""Agente 3 — Generador de salida (sección 7 de ARQUITECTURA.md).

Determinístico, sin LLM. Escribe el `.xlsx` final y corre la verificación
anti-fuga FINAL sobre el archivo ya escrito en disco -- la última línea de
defensa antes de que el archivo se considere entregable. No recibe nunca la
ruta del mapa de auditoría (sección 9): los identificadores que necesita
para su propia verificación llegan ya cargados desde `registros_entrada`,
que el orquestador arma a partir del archivo de entrada original.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

import config
from core.contratos import ENCABEZADOS_SALIDA, ClasificacionClinica, RegistroHistoriaClinica
from core.verificacion_antifuga import ResultadoAntifuga, verificar_archivo_salida

_FUENTE = "Arial"


class ErrorAntifugaSalida(Exception):
    """La verificación anti-fuga final detectó un residuo en el archivo ya
    escrito en disco. El archivo NO debe considerarse entregable -- por
    eso `generar_archivo_salida` lo borra antes de propagar este error."""


def _escribir_workbook(clasificaciones: list[ClasificacionClinica], ruta_destino: Path) -> None:
    ruta_destino.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.title = config.HOJA_SALIDA

    ws.append(ENCABEZADOS_SALIDA)
    for celda in ws[1]:
        celda.font = Font(name=_FUENTE, bold=True)

    for clasificacion in sorted(clasificaciones, key=lambda c: c.codigo_paciente):
        ws.append(
            [
                clasificacion.codigo_paciente,
                clasificacion.enfermedad_probable,
                clasificacion.nivel_riesgo,
                clasificacion.comportamiento_observado,
                clasificacion.confianza,
                clasificacion.antecedentes,
                clasificacion.fuma,
                clasificacion.consumo_sustancias,
            ]
        )

    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name=_FUENTE)

    for columna_celdas in ws.columns:
        longitud = max((len(str(c.value)) for c in columna_celdas if c.value is not None), default=10)
        ws.column_dimensions[columna_celdas[0].column_letter].width = min(max(longitud + 2, 12), 60)

    wb.save(ruta_destino)


def generar_archivo_salida(
    clasificaciones: list[ClasificacionClinica],
    registros_entrada: list[RegistroHistoriaClinica],
    ruta_destino: Path = config.ARCHIVO_SALIDA,
) -> ResultadoAntifuga:
    """Escribe `Formulario_Salida` (mismos 8 encabezados y orden exacto que
    el archivo de referencia, sección 4.3) y corre la verificación
    anti-fuga FINAL sobre el archivo ya escrito en disco (sección 7, punto
    4). Si detecta un residuo, borra el archivo y levanta
    `ErrorAntifugaSalida`: nunca se deja en `data/output/` un entregable
    que no pasó la última línea de defensa.
    """
    _escribir_workbook(clasificaciones, ruta_destino)

    resultado = verificar_archivo_salida(ruta_destino, registros_entrada)
    if not resultado.ok:
        ruta_destino.unlink(missing_ok=True)
        raise ErrorAntifugaSalida(
            f"Residuos detectados en la verificación anti-fuga final: {resultado.categorias_filtradas}"
        )

    return resultado
