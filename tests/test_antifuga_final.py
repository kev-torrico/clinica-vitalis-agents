"""Fase 3 — gate: el archivo `.xlsx` de salida, releído desde disco, no
contiene ninguna cadena identificadora (sección 10.2, "Anti-fuga final").

Incluye también la prueba estática de aislamiento del mapa de auditoría
(sección 10.2: "ni agent2_clasificador.py ni agent3_generador_salida.py
importan o referencian el módulo/ruta del mapa de auditoría")."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

import config
from agents.agente1_anonimizador import cargar_registros_entrada
from agents.agente3_generador_salida import ErrorAntifugaSalida, generar_archivo_salida
from core.contratos import ClasificacionClinica, RegistroHistoriaClinica
from core.verificacion_antifuga import verificar_archivo_salida


def _registro(**overrides) -> RegistroHistoriaClinica:
    base = dict(
        id_registro="REG-99",
        nombre_completo="Ana María Rodríguez Soto",
        cedula="99.888.777",
        direccion="Cra 1 #2-03, Bogotá",
        telefono="300-000-9999",
        fecha_consulta="2026-01-01",
        nota_clinica="texto original irrelevante para esta prueba",
    )
    base.update(overrides)
    return RegistroHistoriaClinica(**base)


def _clasificacion(**overrides) -> ClasificacionClinica:
    base = dict(
        codigo_paciente="PAC-001",
        enfermedad_probable="diabetes tipo 2",
        nivel_riesgo="medio",
        comportamiento_observado="fatiga persistente y sed excesiva",
        confianza="alta",
        antecedentes="antecedente familiar de diabetes tipo 2",
        fuma="No",
        consumo_sustancias="No registrado",
    )
    base.update(overrides)
    return ClasificacionClinica(**base)


def test_verificar_archivo_salida_detecta_un_residuo_real(tmp_path):
    """Si por error un identificador terminara en una celda de salida, la
    verificación final debe detectarlo -- probamos esto escribiendo un xlsx
    'contaminado' a mano, sin pasar por el Agente 3."""
    ruta = tmp_path / "salida_contaminada.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.HOJA_SALIDA
    ws.append(["código_paciente", "comportamiento_observado"])
    ws.append(["PAC-001", "Paciente Ana María Rodríguez Soto acude con fatiga."])
    wb.save(ruta)

    registro = _registro()
    resultado = verificar_archivo_salida(ruta, [registro])

    assert not resultado.ok
    assert any("nombre" in c for c in resultado.categorias_filtradas)


def test_verificar_archivo_salida_sin_residuos_pasa(tmp_path):
    ruta = tmp_path / "salida_limpia.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.HOJA_SALIDA
    ws.append(["código_paciente", "comportamiento_observado"])
    ws.append(["PAC-001", "Fatiga persistente y sed excesiva de 3 meses de evolución."])
    wb.save(ruta)

    registro = _registro()
    resultado = verificar_archivo_salida(ruta, [registro])

    assert resultado.ok
    assert resultado.categorias_filtradas == []


def test_generar_archivo_salida_respeta_encabezados_y_orden_de_referencia(tmp_path):
    ruta_destino = tmp_path / "resultado.xlsx"
    clasificaciones = [
        _clasificacion(codigo_paciente="PAC-002", enfermedad_probable="migraña"),
        _clasificacion(codigo_paciente="PAC-001"),
    ]
    registros = [_registro(id_registro="REG-01"), _registro(id_registro="REG-02", cedula="11.222.333")]

    generar_archivo_salida(clasificaciones, registros, ruta_destino=ruta_destino)

    wb = openpyxl.load_workbook(ruta_destino)
    assert wb.sheetnames == [config.HOJA_SALIDA]  # nada de hoja de entrada ni mapa de auditoría

    ws = wb[config.HOJA_SALIDA]
    encabezados = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    assert list(encabezados) == [
        "código_paciente",
        "enfermedad_probable",
        "nivel_riesgo",
        "comportamiento_observado",
        "confianza",
        "antecedentes",
        "fuma",
        "consumo_sustancias",
    ]

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    assert [f[0] for f in filas] == ["PAC-001", "PAC-002"]  # ordenado por código_paciente

    fuente_encabezado = ws["A1"].font
    assert fuente_encabezado.name == "Arial"
    assert fuente_encabezado.bold is True


def test_generar_archivo_salida_borra_el_archivo_si_la_verificacion_final_falla(tmp_path):
    """Si una clasificación filtrara por error un identificador, el archivo
    NUNCA debe quedar en disco como si fuera entregable."""
    ruta_destino = tmp_path / "resultado_contaminado.xlsx"
    registro = _registro()
    clasificacion_contaminada = _clasificacion(
        comportamiento_observado=f"Paciente {registro.nombre_completo} refiere fatiga."
    )

    with pytest.raises(ErrorAntifugaSalida):
        generar_archivo_salida([clasificacion_contaminada], [registro], ruta_destino=ruta_destino)

    assert not ruta_destino.exists()


def test_pipeline_de_agente3_con_datos_reales_pasa_0_residuos(tmp_path):
    """Gate de Fase 3 (sección 10.2): el .xlsx de salida generado con los
    registros reales de entrada pasa la auditoría anti-fuga en disco con 0
    residuos. Usa clasificaciones simuladas (no llama a OpenRouter) -- lo
    que se prueba aquí es la escritura + verificación final, no el LLM."""
    registros_reales = cargar_registros_entrada(config.ARCHIVO_ENTRADA)
    assert len(registros_reales) == 15

    clasificaciones = [
        _clasificacion(
            codigo_paciente=f"PAC-{i:03d}",
            comportamiento_observado=f"Resumen clínico simulado para el paciente PAC-{i:03d}.",
        )
        for i in range(1, len(registros_reales) + 1)
    ]

    ruta_destino = tmp_path / "Formulario_Salida_Vitalis_resultado.xlsx"
    resultado = generar_archivo_salida(clasificaciones, registros_reales, ruta_destino=ruta_destino)

    assert resultado.ok
    assert resultado.categorias_filtradas == []
    assert ruta_destino.exists()

    # Releído desde disco de forma completamente independiente, por si acaso.
    verificacion_independiente = verificar_archivo_salida(ruta_destino, registros_reales)
    assert verificacion_independiente.ok


# --- Aislamiento del mapa de auditoría (sección 9 y 10.2) ---


def _codigo_fuente(ruta_relativa: str) -> str:
    return (config.BASE_DIR / ruta_relativa).read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "ruta_modulo",
    ["agents/agente2_clasificador.py", "agents/agente3_generador_salida.py"],
)
def test_agentes_2_y_3_no_referencian_el_mapa_de_auditoria(ruta_modulo):
    codigo_fuente = _codigo_fuente(ruta_modulo)
    referencias_prohibidas = ["core.auditoria", "guardar_historias_con_codigo", "ARCHIVO_AUDITORIA"]
    for referencia in referencias_prohibidas:
        assert referencia not in codigo_fuente, f"{ruta_modulo} referencia '{referencia}'"
