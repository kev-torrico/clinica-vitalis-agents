"""Fase 4 — gate: el orquestador coordina las 3 etapas end-to-end y ningún
registro problemático (Agente 1 o Agente 2) tumba la ejecución completa
(sección 13, Fase 4). No llama a OpenRouter real -- `agente2.clasificar_paciente`
va mockeado en todos los casos."""
from __future__ import annotations

import openpyxl
import pytest

import config
import orquestador
from core.cliente_openrouter import ErrorClasificacionAgotada
from core.contratos import ClasificacionClinica


def _clasificacion_falsa(payload) -> ClasificacionClinica:
    return ClasificacionClinica(
        codigo_paciente=payload.codigo_paciente,
        enfermedad_probable="diabetes tipo 2",
        nivel_riesgo="medio",
        comportamiento_observado=f"Resumen clínico simulado para {payload.codigo_paciente}.",
        confianza="alta",
        antecedentes="ninguno registrado",
        fuma="No",
        consumo_sustancias="No registrado",
    )


def _escribir_xlsx_entrada(ruta, filas: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.HOJA_ENTRADA
    ws.append(
        [
            "ID_Registro",
            "Nombre completo",
            "Cédula",
            "Dirección",
            "Teléfono",
            "Fecha de consulta",
            "Nota clínica (texto libre)",
        ]
    )
    for fila in filas:
        ws.append(fila)
    wb.save(ruta)


def test_ejecutar_pipeline_camino_feliz_con_datos_reales(tmp_path, monkeypatch):
    monkeypatch.setattr(orquestador, "clasificar_paciente", _clasificacion_falsa)

    resultado = orquestador.ejecutar_pipeline(
        ruta_entrada=config.ARCHIVO_ENTRADA,
        ruta_auditoria=tmp_path / "auditoria.xlsx",
        ruta_salida=tmp_path / "resultado.xlsx",
    )

    assert len(resultado.clasificaciones) == 15
    assert resultado.requieren_revision_manual == []
    assert resultado.errores_clasificacion == []
    assert resultado.verificacion_final.ok
    assert resultado.ruta_salida.exists()

    wb = openpyxl.load_workbook(resultado.ruta_salida)
    assert wb.sheetnames == [config.HOJA_SALIDA]
    filas = list(wb[config.HOJA_SALIDA].iter_rows(min_row=2, values_only=True))
    assert len(filas) == 15


def test_registro_con_nota_vacia_no_tumba_la_ejecucion_global(tmp_path, monkeypatch):
    monkeypatch.setattr(orquestador, "clasificar_paciente", _clasificacion_falsa)

    ruta_entrada = tmp_path / "entrada.xlsx"
    _escribir_xlsx_entrada(
        ruta_entrada,
        [
            ("REG-01", "Ana Pérez Gómez", "1.111.111", "Calle 1 #1-01, Bogotá", "300-000-0001", "2026-01-01", ""),
            (
                "REG-02",
                "Luis Torres Ruiz",
                "2.222.222",
                "Calle 2 #2-02, Cali",
                "300-000-0002",
                "2026-01-02",
                "Paciente Luis Torres Ruiz, CC 2.222.222, consulta por dolor de cabeza leve.",
            ),
        ],
    )

    resultado = orquestador.ejecutar_pipeline(
        ruta_entrada=ruta_entrada,
        ruta_auditoria=tmp_path / "auditoria.xlsx",
        ruta_salida=tmp_path / "resultado.xlsx",
    )

    assert len(resultado.requieren_revision_manual) == 1
    assert resultado.requieren_revision_manual[0].codigo_paciente == "PAC-001"
    assert len(resultado.clasificaciones) == 1
    assert resultado.clasificaciones[0].codigo_paciente == "PAC-002"
    assert resultado.verificacion_final.ok


def test_fallo_de_clasificacion_no_tumba_la_ejecucion_global(tmp_path, monkeypatch):
    def _clasificar_con_un_fallo(payload):
        if payload.codigo_paciente == "PAC-002":
            raise ErrorClasificacionAgotada("simulación de fallo persistente de OpenRouter")
        return _clasificacion_falsa(payload)

    monkeypatch.setattr(orquestador, "clasificar_paciente", _clasificar_con_un_fallo)

    ruta_entrada = tmp_path / "entrada.xlsx"
    _escribir_xlsx_entrada(
        ruta_entrada,
        [
            (
                "REG-01",
                "Ana Pérez Gómez",
                "1.111.111",
                "Calle 1 #1-01, Bogotá",
                "300-000-0001",
                "2026-01-01",
                "Paciente Ana Pérez Gómez, CC 1.111.111, consulta por dolor lumbar leve.",
            ),
            (
                "REG-02",
                "Luis Torres Ruiz",
                "2.222.222",
                "Calle 2 #2-02, Cali",
                "300-000-0002",
                "2026-01-02",
                "Paciente Luis Torres Ruiz, CC 2.222.222, consulta por dolor de cabeza leve.",
            ),
        ],
    )

    resultado = orquestador.ejecutar_pipeline(
        ruta_entrada=ruta_entrada,
        ruta_auditoria=tmp_path / "auditoria.xlsx",
        ruta_salida=tmp_path / "resultado.xlsx",
    )

    assert resultado.requieren_revision_manual == []
    assert resultado.errores_clasificacion == ["PAC-002"]
    assert [c.codigo_paciente for c in resultado.clasificaciones] == ["PAC-001"]
    assert resultado.verificacion_final.ok

    wb = openpyxl.load_workbook(resultado.ruta_salida)
    filas = list(wb[config.HOJA_SALIDA].iter_rows(min_row=2, values_only=True))
    assert len(filas) == 1  # PAC-002 nunca se inventa una fila para "rellenar"
